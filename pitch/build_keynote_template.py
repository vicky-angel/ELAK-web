#!/usr/bin/env python3
"""Editable ELAK pitch template. Placeholders only — fill the words in Keynote."""

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)
PAPER = RGBColor(0xF4, 0xEF, 0xE4)
INK = RGBColor(0x1D, 0x1D, 0x1F)
MUTED = RGBColor(0x6E, 0x67, 0x5E)
LEAF = RGBColor(0x5F, 0x7D, 0x62)
FONT = "Helvetica Neue"
FONT_EA = "PingFang SC"
LOGO = "/Users/wangbeier/College/Generation AI北大斯坦福中心/Stanford Hackathon/ELAK-Physicalcare-therapy/pitch/assets/elak-logo.png"
OUT = "/Users/wangbeier/College/Generation AI北大斯坦福中心/Stanford Hackathon/ELAK-Physicalcare-therapy/pitch/ELAK Pitch.pptx"

TOTAL = 8


def set_run_font(run, size, color, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", FONT), ("a:ea", FONT_EA), ("a:cs", FONT)):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", face)


def tracking(run, hundredths):
    run._r.get_or_add_rPr().set("spc", str(hundredths))


def add_text(slide, l, t, w, h, text, size, color, *, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, track=None):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Emu(0)
    tf.margin_right = Emu(0)
    tf.margin_top = Emu(0)
    tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    lines = text.split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        run = p.add_run()
        run.text = line
        set_run_font(run, size, color, bold)
        if track:
            tracking(run, track)
    return box


def paint(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER


def chrome(slide, n, section):
    paint(slide)
    slide.shapes.add_picture(LOGO, Inches(0.58), Inches(0.26), height=Inches(0.78))
    if section:
        add_text(slide, 1.55, 0.48, 6.5, 0.36, section, 13, LEAF, track=280)
    add_text(
        slide, 10.3, 0.48, 2.35, 0.36,
        f"{n:02d}  /  {TOTAL:02d}",
        13, MUTED, align=PP_ALIGN.RIGHT, track=180,
    )


def headline(slide, text, top=2.15, size=48):
    add_text(slide, 0.72, top, 11.8, 1.7, text, size, INK)


def subline(slide, text, top=4.15):
    add_text(slide, 0.72, top, 10.5, 1.3, text, 24, MUTED)


def columns(slide, items, top=4.15):
    gap = 0.38
    width = 3.72
    x = 0.72
    for label, body in items:
        add_text(slide, x, top, width, 0.36, label, 13, LEAF, track=220)
        add_text(slide, x, top + 0.48, width, 1.7, body, 22, INK)
        x += width + gap


def rows(slide, pairs, top=3.85):
    y = top
    for left, right in pairs:
        add_text(slide, 0.72, y, 3.3, 0.55, left, 24, INK)
        add_text(slide, 4.2, y, 8.2, 0.55, right, 24, MUTED)
        y += 0.72


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    prs.core_properties.title = "ELAK Pitch"
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    chrome(s, 1, "")
    headline(s, "Title", top=2.35, size=64)
    subline(s, "One line about what you do", top=3.85)
    add_text(s, 0.72, 6.55, 8, 0.4, "Your names    ·    Date", 16, MUTED, track=80)

    # 2 Problem
    s = prs.slides.add_slide(blank)
    chrome(s, 2, "PROBLEM")
    headline(s, "Headline")
    columns(s, [
        ("NOW", "What people do today"),
        ("WHY IT FAILS", "Why that is not enough"),
        ("SUBSTITUTES", "Why the other options fail"),
    ])

    # 3 Solution
    s = prs.slides.add_slide(blank)
    chrome(s, 3, "SOLUTION")
    headline(s, "Headline")
    columns(s, [
        ("TODAY", "The current step"),
        ("WITH YOU", "What changes"),
        ("AFTER", "The result"),
    ])

    # 4 Technology
    s = prs.slides.add_slide(blank)
    chrome(s, 4, "TECHNOLOGY")
    headline(s, "Headline")
    columns(s, [
        ("HOW IT WORKS", "The product, in one or two lines"),
        ("WHAT IS NEW", "What you can show"),
        ("DATA", "What comes back"),
    ])

    # 5 Competition
    s = prs.slides.add_slide(blank)
    chrome(s, 5, "COMPETITION")
    headline(s, "Headline", top=1.9, size=44)
    rows(s, [
        ("You", "Where you are strong"),
        ("Them", "Where they stop"),
        ("Them", "Where they stop"),
        ("Them", "Where they stop"),
    ], top=3.9)

    # 6 Business model
    s = prs.slides.add_slide(blank)
    chrome(s, 6, "BUSINESS MODEL")
    headline(s, "Headline", top=1.9, size=44)
    rows(s, [
        ("Who pays", "The buyer"),
        ("Price", "How you charge"),
        ("First customer", "Who, and why they start"),
        ("Channel", "How you reach them"),
    ], top=3.9)

    # 7 Team
    s = prs.slides.add_slide(blank)
    chrome(s, 7, "TEAM")
    headline(s, "Headline", top=1.9, size=44)
    rows(s, [
        ("Name", "Role"),
        ("Name", "Role"),
        ("Name", "Role"),
        ("Advisor", "Role"),
    ], top=3.9)

    # 8 Backup
    s = prs.slides.add_slide(blank)
    chrome(s, 8, "BACKUP")
    headline(s, "Headline")
    columns(s, [
        ("MARKET", "The number, when you have it"),
        ("ROADMAP", "The next 90 days"),
        ("EVIDENCE", "What a customer already said"),
    ])

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
